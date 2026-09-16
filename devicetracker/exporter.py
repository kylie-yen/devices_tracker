"""导出为 CSV / Excel。

用标准库 csv 加 openpyxl，而不是 pandas：pandas 光 import 就要 0.5～1.5 秒，
为一个导出功能拖慢桌面应用的启动不值得；而且 openpyxl 能直接控制列宽、
数字格式和冻结窗格，比 ``df.to_excel`` 之后再用 openpyxl 打开一遍更省事。

导出内容包含派生列（持有天数、两个口径的日均），因为导出本来就是为了看这些。
"""

from __future__ import annotations

import csv
from collections.abc import Iterable, Sequence
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import Device

HEADERS: tuple[str, ...] = (
    "名称",
    "分类",
    "购入日期",
    "购入价格",
    "售出日期",
    "售出价格",
    "持有天数",
    "购入成本日均",
    "净成本日均",
    "状态",
    "备注",
)

# HEADERS 中属于金额的列下标，Excel 里要写成数值再套数字格式
MONEY_COLUMNS: tuple[int, ...] = (3, 5, 7, 8)

COLUMN_WIDTHS: tuple[int, ...] = (24, 10, 12, 12, 12, 12, 10, 14, 14, 10, 28)

_SHEET_TITLE = "设备台账"


def row_for(device: Device, today: date) -> list[str | float | int]:
    """把设备摊平成一行的导出值。

    金额保持 float、天数保持 int，而不是格式化后的字符串 —— Excel 里必须是
    真数值，否则用户没法求和、排序。
    """
    return [
        device.name,
        device.category,
        device.purchase_date.isoformat(),
        device.purchase_price,
        device.sale_date.isoformat() if device.sale_date else "",
        device.sale_price if device.sale_price is not None else "",
        device.days_held(today),
        round(device.daily_purchase_cost(today), 2),
        round(device.daily_net_cost(today), 2),
        device.status_text,
        device.note,
    ]


def rows_for(devices: Iterable[Device], today: date) -> list[list]:
    return [row_for(device, today) for device in devices]


def export_csv(devices: Sequence[Device], path: Path | str, today: date) -> Path:
    """写出 CSV。

    ``encoding="utf-8-sig"`` 是为了带 BOM —— 否则 Excel 双击打开中文是乱码。
    ``newline=""`` 是 csv 模块的要求，缺了它 Windows 上每行之间会多一个空行。
    """
    path = Path(path)
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(rows_for(devices, today))
    return path


def export_excel(devices: Sequence[Device], path: Path | str, today: date) -> Path:
    """写出 xlsx，带表头加粗、列宽和冻结首行。"""
    path = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _SHEET_TITLE

    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in rows_for(devices, today):
        sheet.append(row)

    # 金额列套数字格式，保证在 Excel 里能直接求和
    for row in sheet.iter_rows(min_row=2, max_col=len(HEADERS)):
        for index in MONEY_COLUMNS:
            row[index].number_format = "#,##0.00"

    sheet.freeze_panes = "A2"
    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(path)
    return path
